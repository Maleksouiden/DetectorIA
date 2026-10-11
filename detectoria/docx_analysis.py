"""Extraction du texte d'un document Word, découpage et surlignage."""

from __future__ import annotations

import re
import threading
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

from detectoria.pdf_analysis import french_ratio, make_chunks, verdict

REFORM_HEADING = "Reformulations proposées — Détecteur IA"
W_R = qn("w:r")
W_T = qn("w:t")
W_HL = qn("w:highlight")
W_RPR = qn("w:rPr")


def analyse_docx(path, detector, options: dict, out_path, progress=lambda *a: None, cancel: threading.Event | None = None):
    doc = Document(path)
    paragraphs = list(_iter_paragraphs(doc))
    units = []
    sample_words: list[str] = []
    for number, paragraph in enumerate(paragraphs, start=1):
        words = _paragraph_words(paragraph, number, len(units))
        if not _keep_paragraph(words):
            continue
        if len(sample_words) < 4000:
            sample_words.extend(word["text"] for word in words[:400])
        units.append(words)

    flat = [word for words in units for word in words]
    spans = make_chunks(
        flat,
        int(options["min_words"]),
        int(options["target_words"]),
        int(options["max_words"]),
    )
    threshold = float(options["threshold"])
    strong = float(options["strong_prob"])
    findings = []
    passages = []
    analysed_words = flagged_words = 0
    weighted = 0.0
    stopped = False
    highlights: dict[int, list[tuple]] = {}

    batch = max(1, int(options.get("batch_size") or 8))
    total = max(len(spans), 1)
    for offset in range(0, len(spans), batch):
        if cancel is not None and cancel.is_set():
            stopped = True
            break
        chunk = spans[offset : offset + batch]
        progress(offset, total, f"Analyse des passages {offset + 1}–{offset + len(chunk)}/{len(spans)}…")
        texts = [" ".join(word["text"] for word in flat[start:end]) for start, end in chunk]
        probs = detector.predict(texts)
        for (start, end), text, prob in zip(chunk, texts, probs):
            count = end - start
            analysed_words += count
            weighted += prob * count
            passage = {
                "page": flat[start]["para_no"],
                "prob": prob,
                "words": count,
                "text": text,
                "reformulation": "",
                "flagged": prob >= threshold,
            }
            passages.append(passage)
            if prob < threshold:
                continue
            flagged_words += count
            findings.append(passage)
            color = "red" if prob >= strong else "yellow"
            for group in _same_paragraph(flat[start:end]):
                key = id(group[0]["para"])
                highlights.setdefault(key, []).append((group[0]["start"], group[-1]["end"], color, group[0]["para"]))

    if analysed_words and not stopped:
        for ranges in highlights.values():
            paragraph = ranges[0][3]
            _highlight_paragraph(paragraph, [(start, end, color) for start, end, color, _para in ranges])
        doc.save(out_path)
    progress(total, total, "Terminé" if not stopped else "Analyse interrompue")

    share = flagged_words / analysed_words if analysed_words else 0.0
    return {
        "pages": len(paragraphs),
        "analysed_words": analysed_words,
        "flagged_words": flagged_words,
        "share": share,
        "mean_prob": weighted / analysed_words if analysed_words else 0.0,
        "verdict": verdict(share) if analysed_words else "Aucun texte exploitable",
        "findings": findings,
        "passages": passages,
        "french_ratio": french_ratio(sample_words),
        "out_path": str(out_path) if analysed_words and not stopped else None,
        "cancelled": stopped,
        "model_label": getattr(detector, "spec", {}).get("label", ""),
        "calibration": getattr(detector, "calibration", ""),
        "kind": "docx",
        "place_label": "§",
    }


def write_docx_reformulations(path: str, findings: list[dict]) -> None:
    if not path or not findings:
        return
    doc = Document(path)
    _strip_reform_appendix(doc)
    written = [item for item in findings if (item.get("reformulation") or "").strip()]
    if written:
        doc.add_page_break()
        doc.add_heading(REFORM_HEADING, level=1)
        note = doc.add_paragraph(
            "Ces formulations sont des propositions. Elles ne prouvent pas qu’un texte est humain."
        )
        if note.runs:
            note.runs[0].italic = True
        for item in written:
            title = doc.add_paragraph()
            run = title.add_run(f"Paragraphe {item['page']} · {item['prob']:.0%}")
            run.bold = True
            doc.add_paragraph((item.get("reformulation") or "").strip())
    doc.save(path)


def _iter_paragraphs(doc):
    seen_cells: set[int] = set()

    def walk_table(table):
        for row in table.rows:
            for cell in row.cells:
                key = id(cell._tc)
                if key in seen_cells:
                    continue
                seen_cells.add(key)
                yield from cell.paragraphs
                for nested in cell.tables:
                    yield from walk_table(nested)

    yield from doc.paragraphs
    for table in doc.tables:
        yield from walk_table(table)


def _run_elements(paragraph: Paragraph):
    for child in paragraph._p:
        if child.tag == W_R:
            yield child
        elif child.tag == qn("w:hyperlink"):
            for node in child:
                if node.tag == W_R:
                    yield node


def _run_text(run) -> str:
    return "".join(node.text or "" for node in run.findall(W_T))


def _plain_text(paragraph: Paragraph) -> str:
    return "".join(_run_text(run) for run in _run_elements(paragraph))


def _paragraph_words(paragraph: Paragraph, para_no: int, line_id: int) -> list[dict]:
    text = _plain_text(paragraph)
    words = []
    for match in re.finditer(r"\S+", text):
        words.append({
            "text": match.group(),
            "start": match.start(),
            "end": match.end(),
            "para": paragraph,
            "para_no": para_no,
            "line": (line_id, 0),
        })
    return words


def _keep_paragraph(words: list[dict]) -> bool:
    if len(words) < 5:
        return False
    text = " ".join(word["text"] for word in words)
    letters = sum(ch.isalpha() for ch in text)
    if letters < 0.6 * max(len(text), 1):
        return False
    if re.search(r"\.{4,}|(\. ){4,}", text):
        return False
    return True


def _same_paragraph(words: list[dict]):
    group: list[dict] = []
    current = None
    for word in words:
        key = id(word["para"])
        if current is None:
            current = key
        if key != current:
            yield group
            group = [word]
            current = key
        else:
            group.append(word)
    if group:
        yield group


def _highlight_paragraph(paragraph: Paragraph, spans: list[tuple[int, int, str]]) -> None:
    if not spans:
        return
    spans = sorted(spans)
    pos = 0
    for run in list(_run_elements(paragraph)):
        text = _run_text(run)
        length = len(text)
        local = []
        for start, end, color in spans:
            if end <= pos or start >= pos + length:
                continue
            local.append((max(start, pos) - pos, min(end, pos + length) - pos, color))
        if local:
            _split_run(run, local)
        pos += length


def _split_run(run, local: list[tuple[int, int, str]]) -> None:
    text = _run_text(run)
    if not text:
        return
    colors: list[str | None] = [None] * len(text)
    for start, end, color in local:
        for index in range(max(0, start), min(len(text), end)):
            colors[index] = color
    if all(color is None for color in colors):
        return
    segments: list[tuple[str, str | None]] = []
    index = 0
    while index < len(text):
        nxt = index + 1
        while nxt < len(text) and colors[nxt] == colors[index]:
            nxt += 1
        segments.append((text[index:nxt], colors[index]))
        index = nxt
    _set_run_text(run, segments[0][0])
    _set_highlight(run, segments[0][1])
    anchor = run
    for seg_text, color in segments[1:]:
        clone = deepcopy(run)
        _set_run_text(clone, seg_text)
        _set_highlight(clone, color)
        anchor.addnext(clone)
        anchor = clone


def _set_run_text(run, text: str) -> None:
    for node in list(run.findall(W_T)):
        run.remove(node)
    node = OxmlElement("w:t")
    if text[:1].isspace() or text[-1:].isspace():
        node.set(qn("xml:space"), "preserve")
    node.text = text
    run.append(node)


def _set_highlight(run, color: str | None) -> None:
    props = run.find(W_RPR)
    if color is None:
        if props is not None:
            mark = props.find(W_HL)
            if mark is not None:
                props.remove(mark)
        return
    if props is None:
        props = OxmlElement("w:rPr")
        run.insert(0, props)
    mark = props.find(W_HL)
    if mark is None:
        mark = OxmlElement("w:highlight")
        props.append(mark)
    mark.set(qn("w:val"), color)


def _is_page_break(element) -> bool:
    if element.tag != qn("w:p"):
        return False
    if "".join(node.text or "" for node in element.iter(W_T)).strip():
        return False
    return element.find(".//" + qn("w:br")) is not None


def _strip_reform_appendix(doc) -> None:
    body = doc.element.body
    children = list(body)
    cut = None
    for index, child in enumerate(children):
        if child.tag != qn("w:p"):
            continue
        if REFORM_HEADING in "".join(node.text or "" for node in child.iter(W_T)):
            cut = index
            break
    if cut is None:
        return
    if cut > 0 and _is_page_break(children[cut - 1]):
        cut -= 1
    for child in children[cut:]:
        if child.tag == qn("w:sectPr"):
            continue
        body.remove(child)
